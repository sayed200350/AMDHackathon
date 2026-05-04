"""Export utilities — converts memo Markdown to DOCX format."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass


@dataclass(frozen=True)
class ExportResult:
    """Result of a memo export operation."""

    content: bytes
    filename: str
    content_type: str


def memo_to_docx(memo_markdown: str, matter_name: str = "Contract Review") -> ExportResult:
    """Convert a Markdown memo to DOCX format.

    Args:
        memo_markdown: The memo content in Markdown format.
        matter_name: Name for the output file.

    Returns:
        ExportResult with the DOCX bytes, filename, and content type.
    """
    from docx import Document
    from docx.shared import Pt

    doc = Document()

    # Style defaults
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    for line in memo_markdown.split("\n"):
        stripped = line.strip()

        if not stripped:
            doc.add_paragraph("")
            continue

        # Headings
        if stripped.startswith("# ") and not stripped.startswith("## "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("---"):
            doc.add_paragraph("_" * 50)
        elif stripped.startswith("| "):
            # Simple table row — add as formatted paragraph
            cells = [c.strip() for c in stripped.split("|") if c.strip()]
            if cells and not all(c.startswith("-") for c in cells):
                para = doc.add_paragraph()
                para.style = doc.styles["Normal"]
                run = para.add_run("  |  ".join(cells))
                run.font.size = Pt(10)
                run.font.name = "Consolas"
        elif stripped.startswith("- "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        elif re.match(r"^\d+\.", stripped):
            text = re.sub(r"^\d+\.\s*", "", stripped)
            doc.add_paragraph(text, style="List Number")
        elif stripped.startswith("**") and stripped.endswith("**"):
            para = doc.add_paragraph()
            run = para.add_run(stripped.strip("*"))
            run.bold = True
        else:
            # Handle inline bold markers
            para = doc.add_paragraph()
            parts = re.split(r"(\*\*.*?\*\*)", stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = para.add_run(part.strip("*"))
                    run.bold = True
                else:
                    para.add_run(part)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    safe_name = re.sub(r"[^\w\s-]", "", matter_name).strip().replace(" ", "_")

    return ExportResult(
        content=buffer.read(),
        filename=f"{safe_name}_Review_Memo.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
